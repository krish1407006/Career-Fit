import io
import re

from django.core.files.uploadedfile import UploadedFile

from pypdf import PdfReader


def _extract_pdf(file_obj):
    reader = PdfReader(file_obj)
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages)


def _extract_docx(file_obj):
    from docx import Document

    doc = Document(file_obj)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_txt(file_obj):
    raw = file_obj.read()
    if isinstance(raw, bytes):
        raw = raw.decode("utf-8", errors="replace")
    return raw or ""


def _clean_text(text):
    text = re.sub(r"\r\n?", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text_from_resume(file_obj):
    """Extract plain text from an uploaded resume (pdf / docx / txt)."""
    if not isinstance(file_obj, UploadedFile):
        data = file_obj.read()
        name = getattr(file_obj, "name", "")
        file_obj = UploadedFile(io.BytesIO(data), name=name)
        file_obj.seek(0)

    name = (file_obj.name or "").lower()
    if name.endswith(".pdf"):
        return _clean_text(_extract_pdf(file_obj))
    if name.endswith(".docx"):
        return _clean_text(_extract_docx(file_obj))
    if name.endswith(".txt"):
        return _clean_text(_extract_txt(file_obj))
    raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT resume.")